"""
PPTX Extractor Module
Extracts text and structure from PowerPoint documents using markitdown.
"""

import subprocess
import re
from typing import List, Dict, Any, Optional
from dataclasses import dataclass, asdict


@dataclass
class Section:
    """Represents a document section with heading and content."""
    heading: str
    content: str
    hierarchy: int


class PPTXExtractor:
    """
    Extracts text and structured content from PPTX documents.
    Uses markitdown for markdown conversion.
    """

    def __init__(self):
        self.has_markitdown = self._check_markitdown()
        self.pptx_module = None

    def _check_markitdown(self) -> bool:
        """Check if markitdown is available."""
        try:
            result = subprocess.run(
                ['python', '-m', 'markitdown', '--help'],
                capture_output=True,
                text=True,
                timeout=10
            )
            return result.returncode == 0
        except (FileNotFoundError, subprocess.TimeoutExpired):
            # Try importing as a Python module
            try:
                from markitdown import MarkItDown
                return True
            except ImportError:
                return False

    def _import_pptx(self):
        """Lazy import of python-pptx."""
        try:
            from pptx import Presentation
            self.pptx_module = Presentation
        except ImportError:
            raise ImportError(
                "python-pptx is required for PPTX extraction fallback. "
                "Install it with: pip install python-pptx"
            )

    def extract(self, file_path: str) -> Dict[str, Any]:
        """
        Extract text and structure from a PPTX file.

        Args:
            file_path: Path to the PPTX file

        Returns:
            Dictionary containing extracted text and structured sections
        """
        # Try markitdown first
        if self.has_markitdown:
            try:
                return self._extract_with_markitdown(file_path)
            except Exception:
                pass  # Fall through to fallback

        # Fallback to python-pptx
        return self._extract_with_python_pptx(file_path)

    def _extract_with_markitdown(self, file_path: str) -> Dict[str, Any]:
        """
        Extract using markitdown (converts PPTX to markdown).

        Args:
            file_path: Path to the PPTX file

        Returns:
            Extraction result dictionary
        """
        try:
            # Try using markitdown as Python module first
            try:
                from markitdown import MarkItDown
                md = MarkItDown()
                result = md.convert(file_path)
                markdown_text = result.text_content
            except ImportError:
                # Fall back to CLI
                result = subprocess.run(
                    ['python', '-m', 'markitdown', file_path],
                    capture_output=True,
                    text=True,
                    timeout=60
                )

                if result.returncode != 0:
                    raise ValueError(f"Markitdown error: {result.stderr}")

                markdown_text = result.stdout

            # Parse sections from markdown
            sections = self._parse_markdown_sections(markdown_text)

            # Convert to plain text
            plain_text = self._markdown_to_plain(markdown_text)

            # Count slides from section structure
            slide_count = self._count_slides(sections)

            return {
                'extractedText': plain_text,
                'structuredSections': [asdict(s) for s in sections],
                'metadata': {
                    'pageCount': slide_count,
                    'wordCount': len(plain_text.split()),
                    'extractionMethod': 'markitdown'
                }
            }

        except subprocess.TimeoutExpired:
            raise ValueError(f"Markitdown timed out processing: {file_path}")
        except Exception as e:
            raise ValueError(f"Failed to extract PPTX with markitdown: {str(e)}")

    def _extract_with_python_pptx(self, file_path: str) -> Dict[str, Any]:
        """
        Fallback extraction using python-pptx.

        Args:
            file_path: Path to the PPTX file

        Returns:
            Extraction result dictionary
        """
        if not self.pptx_module:
            self._import_pptx()

        try:
            prs = self.pptx_module(file_path)
            sections = []
            all_text = []

            for slide_num, slide in enumerate(prs.slides, 1):
                slide_title = f"Slide {slide_num}"
                slide_content = []

                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text = shape.text.strip()

                        # First text shape is often the title
                        if not slide_content and len(text) < 100:
                            slide_title = text
                        else:
                            slide_content.append(text)

                        all_text.append(text)

                    # Handle tables
                    if shape.has_table:
                        table_text = self._extract_table(shape.table)
                        if table_text:
                            slide_content.append(table_text)
                            all_text.append(table_text)

                sections.append(Section(
                    heading=slide_title,
                    content="\n".join(slide_content),
                    hierarchy=1
                ))

            full_text = "\n\n".join(all_text)

            return {
                'extractedText': full_text,
                'structuredSections': [asdict(s) for s in sections],
                'metadata': {
                    'pageCount': len(prs.slides),
                    'wordCount': len(full_text.split()),
                    'extractionMethod': 'python-pptx'
                }
            }

        except Exception as e:
            if "password" in str(e).lower() or "encrypted" in str(e).lower():
                raise ValueError(f"PPTX is password-protected: {file_path}")
            raise ValueError(f"Failed to extract PPTX: {str(e)}")

    def _extract_table(self, table) -> str:
        """Extract text from a PowerPoint table."""
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        return "\n".join(rows)

    def _parse_markdown_sections(self, markdown: str) -> List[Section]:
        """
        Parse markdown content into structured sections (slides).

        Args:
            markdown: Markdown text from markitdown

        Returns:
            List of Section objects
        """
        sections = []
        lines = markdown.split('\n')
        current_heading = "Slide 1"
        current_content = []
        current_hierarchy = 1
        slide_num = 1

        for line in lines:
            # Check for slide markers (headings or horizontal rules)
            heading_match = re.match(r'^(#{1,6})\s+(.+)$', line)
            slide_break = line.strip() in ['---', '***', '___']

            if heading_match or slide_break:
                # Save previous section
                if current_content or heading_match:
                    if current_content:
                        content_text = "\n".join(current_content).strip()
                        content_text = self._clean_markdown(content_text)
                        sections.append(Section(
                            heading=current_heading,
                            content=content_text,
                            hierarchy=current_hierarchy
                        ))

                if heading_match:
                    current_hierarchy = len(heading_match.group(1))
                    current_heading = heading_match.group(2).strip()
                elif slide_break:
                    slide_num += 1
                    current_heading = f"Slide {slide_num}"
                    current_hierarchy = 1

                current_content = []
            else:
                if line.strip():
                    current_content.append(line)

        # Save last section
        if current_content:
            content_text = "\n".join(current_content).strip()
            content_text = self._clean_markdown(content_text)
            sections.append(Section(
                heading=current_heading,
                content=content_text,
                hierarchy=current_hierarchy
            ))

        # If no sections found, create one with all content
        if not sections and markdown.strip():
            sections.append(Section(
                heading="Presentation Content",
                content=self._clean_markdown(markdown),
                hierarchy=1
            ))

        return sections

    def _count_slides(self, sections: List[Section]) -> int:
        """Estimate slide count from sections."""
        slide_count = 0
        for section in sections:
            if section.heading.lower().startswith('slide') or section.hierarchy == 1:
                slide_count += 1
        return max(1, slide_count)

    def _markdown_to_plain(self, markdown: str) -> str:
        """Convert markdown to plain text."""
        text = markdown

        # Remove heading markers
        text = re.sub(r'^#{1,6}\s+', '', text, flags=re.MULTILINE)

        # Remove bold/italic markers
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        text = re.sub(r'\*(.+?)\*', r'\1', text)

        # Remove link formatting
        text = re.sub(r'\[(.+?)\]\(.+?\)', r'\1', text)

        # Remove horizontal rules
        text = re.sub(r'^[-*_]{3,}\s*$', '', text, flags=re.MULTILINE)

        # Remove image references
        text = re.sub(r'!\[.*?\]\(.*?\)', '', text)

        return text.strip()

    def _clean_markdown(self, text: str) -> str:
        """Remove common markdown artifacts from text."""
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        text = re.sub(r'\*(.+?)\*', r'\1', text)
        text = re.sub(r'\[(.+?)\]\(.+?\)', r'\1', text)
        text = re.sub(r'\n{3,}', '\n\n', text)
        return text.strip()

    def can_handle(self, file_path: str) -> bool:
        """Check if this extractor can handle the given file."""
        return file_path.lower().endswith('.pptx')
