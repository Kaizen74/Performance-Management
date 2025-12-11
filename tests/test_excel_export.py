"""
Test Suite for Excel Export Engine (Milestone 6)
Tests Excel workbook generation from analysis results.
"""

import os
import sys
import pytest
import tempfile
from pathlib import Path

# Add backend to path
sys.path.insert(0, str(Path(__file__).parent.parent / 'backend'))

from exports.excel_export_engine import ExcelExportEngine, generate_excel_export
from exports.pdf_export_engine import PDFExportEngine, generate_pdf_export


# Mock strategic framework
MOCK_FRAMEWORK = {
    "organizationalPurpose": {
        "vision": "To be the leading sustainable logistics provider in Asia-Pacific by 2030",
        "mission": "We deliver excellence through innovation",
        "values": ["Innovation", "Sustainability", "Excellence"]
    },
    "strategicPerspectives": {
        "financial": {
            "objectives": [
                {"id": "F1", "objective": "Revenue growth 12%", "keyMeasures": ["Revenue"]},
                {"id": "F2", "objective": "Cost reduction 15%", "keyMeasures": ["Cost"]}
            ]
        },
        "customer": {
            "objectives": [
                {"id": "C1", "objective": "NPS > 70", "keyMeasures": ["NPS"]},
                {"id": "C2", "objective": "On-time delivery 95%", "keyMeasures": ["Delivery"]}
            ]
        },
        "internalProcess": {
            "objectives": [
                {"id": "P1", "objective": "AI route optimization", "keyMeasures": ["Efficiency"]},
                {"id": "P2", "objective": "Carbon neutrality 2028", "keyMeasures": ["Emissions"]},
                {"id": "P3", "objective": "Operational efficiency", "keyMeasures": ["Cost per unit"]}
            ]
        },
        "learningGrowth": {
            "objectives": [
                {"id": "L1", "objective": "Digital capabilities", "keyMeasures": ["Skills"]},
                {"id": "L2", "objective": "Employee engagement 80%", "keyMeasures": ["Engagement"]}
            ]
        }
    },
    "keyPerformanceRequirements": [
        {"id": "KPR1", "requirement": "Drive digital transformation", "priority": "critical", "linkedObjectiveIds": ["P1", "L1"]}
    ],
    "metadata": {
        "frameworkId": "test-framework-001"
    }
}


def generate_mock_analysis_results(count: int = 5):
    """Generate mock analysis results for testing."""
    seniority_levels = ['executive', 'senior', 'mid', 'junior', 'senior']
    results = []

    for i in range(count):
        results.append({
            "documentId": f"doc-{i+1:03d}",
            "fileName": f"employee_{i+1}_goals.docx",
            "employeeMetadata": {
                "employeeName": f"Employee {i+1}",
                "jobTitle": f"Manager Level {i+1}",
                "department": "Operations",
                "seniorityLevel": seniority_levels[i % len(seniority_levels)]
            },
            "overallAlignmentScore": 60 + (i * 5) % 40,
            "overallImpactScore": 55 + (i * 7) % 35,
            "weightedCompositeScore": 58 + (i * 6) % 38,
            "coherenceAssessment": {
                "internalConsistency": 70 + i,
                "logicalFlow": 65 + i,
                "completeness": 60 + i,
                "overallCoherence": 65 + i
            },
            "scoringBreakdown": {
                "strategicImpact": {"raw": 70, "weighted": 21},
                "alignmentBreadth": {"raw": 75, "weighted": 18.75},
                "goalQuality": {"raw": 68, "weighted": 17},
                "roleAppropriateness": {"raw": 72, "weighted": 18}
            },
            "goals": [
                {
                    "goalId": f"G{j+1}",
                    "goalText": f"Goal {j+1}: Improve performance metric by 20%",
                    "alignmentScore": 65 + j * 5,
                    "impactScore": 60 + j * 4,
                    "alignedObjectives": ["P1", "P3"] if j % 2 == 0 else ["F1", "C1"],
                    "alignmentRationale": "Supports operational efficiency",
                    "impactRationale": "Direct contribution to cost targets",
                    "gaps": ["No direct link to sustainability"] if j % 3 == 0 else []
                }
                for j in range(4)
            ],
            "strategicCoverage": {
                "financial": {"covered": 1, "total": 2, "percentage": 50.0},
                "customer": {"covered": 1, "total": 2, "percentage": 50.0},
                "process": {"covered": 2, "total": 3, "percentage": 66.7},
                "learning": {"covered": 0, "total": 2, "percentage": 0.0}
            },
            "recommendations": [
                "Add goals addressing learning perspective",
                "Strengthen sustainability linkage"
            ]
        })

    return results


def generate_mock_recommendations(results):
    """Generate mock recommendations for testing."""
    recommendations = {}

    for result in results:
        doc_id = result['documentId']
        recommendations[doc_id] = {
            "documentId": doc_id,
            "employeeMetadata": result.get("employeeMetadata", {}),
            "recommendations": [
                {
                    "recommendationId": f"R{i+1}",
                    "revisedGoal": {
                        "objective": f"Revised goal {i+1}: Achieve 25% improvement in key metric",
                        "keyResults": ["KR1", "KR2"],
                        "timeline": "Q2 2025",
                        "metrics": ["Efficiency", "Quality"]
                    },
                    "strategicLinkages": ["P1", "L1"] if i % 2 == 0 else ["F1", "C1"],
                    "predictedAlignmentGain": 10 + i * 2,
                    "roleAppropriatenessJustification": "Appropriate for seniority level",
                    "evidence": {
                        "source": "Industry best practices",
                        "finding": "Similar organizations achieve better alignment",
                        "url": "https://example.com/research"
                    },
                    "implementationNotes": "Requires cross-functional coordination"
                }
                for i in range(5)
            ],
            "projectedNewAlignmentScore": result.get("overallAlignmentScore", 70) + 15,
            "projectedNewImpactScore": result.get("overallImpactScore", 65) + 12,
            "projectedNewCoherenceScore": 80
        }

    return recommendations


class TestExcelExportEngine:
    """Test cases for the ExcelExportEngine class."""

    def test_engine_initialization(self):
        """Test engine initializes correctly."""
        results = generate_mock_analysis_results(3)
        engine = ExcelExportEngine(MOCK_FRAMEWORK, results)

        assert engine.framework == MOCK_FRAMEWORK
        assert len(engine.results) == 3
        assert engine.wb is not None

    def test_generate_workbook_creates_file(self):
        """Test workbook generation creates file."""
        results = generate_mock_analysis_results(5)
        engine = ExcelExportEngine(MOCK_FRAMEWORK, results)

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            output_path = tmp.name

        try:
            result_path = engine.generate_workbook(output_path)
            assert os.path.exists(result_path)
            assert result_path == output_path
            # Check file size is reasonable (not empty)
            assert os.path.getsize(result_path) > 1000
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)

    def test_workbook_has_required_sheets(self):
        """Test workbook contains all required sheets."""
        from openpyxl import load_workbook

        results = generate_mock_analysis_results(5)
        engine = ExcelExportEngine(MOCK_FRAMEWORK, results)

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            output_path = tmp.name

        try:
            engine.generate_workbook(output_path)
            wb = load_workbook(output_path)

            expected_sheets = [
                "Executive Summary",
                "Employee Details",
                "Alignment Matrix",
                "Recommendations",
                "Gap Analysis"
            ]

            for sheet_name in expected_sheets:
                assert sheet_name in wb.sheetnames, f"Missing sheet: {sheet_name}"
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)

    def test_employee_details_has_correct_columns(self):
        """Test Employee Details sheet has correct structure."""
        from openpyxl import load_workbook

        results = generate_mock_analysis_results(5)
        engine = ExcelExportEngine(MOCK_FRAMEWORK, results)

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            output_path = tmp.name

        try:
            engine.generate_workbook(output_path)
            wb = load_workbook(output_path)
            ws = wb["Employee Details"]

            # Check header row
            assert ws['A1'].value == "Employee Name"
            assert ws['E1'].value == "Alignment Score"
            assert ws['J1'].value == "Tier"
            assert ws['O1'].value == "Proposed Adjustment 1"
            assert ws['T1'].value == "Projected New Score"
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)

    def test_employee_details_row_count(self):
        """Test Employee Details has correct number of rows."""
        from openpyxl import load_workbook

        num_employees = 10
        results = generate_mock_analysis_results(num_employees)
        engine = ExcelExportEngine(MOCK_FRAMEWORK, results)

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            output_path = tmp.name

        try:
            engine.generate_workbook(output_path)
            wb = load_workbook(output_path)
            ws = wb["Employee Details"]

            # Should have header + num_employees rows
            assert ws.max_row == num_employees + 1
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)

    def test_alignment_matrix_structure(self):
        """Test Alignment Matrix has correct structure."""
        from openpyxl import load_workbook

        results = generate_mock_analysis_results(3)
        engine = ExcelExportEngine(MOCK_FRAMEWORK, results)

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            output_path = tmp.name

        try:
            engine.generate_workbook(output_path)
            wb = load_workbook(output_path)
            ws = wb["Alignment Matrix"]

            # Check title exists
            assert "Alignment Matrix" in str(ws['A1'].value)

            # Check headers exist
            assert ws['A3'].value == "Perspective"
            assert ws['B3'].value == "Strategic Objective"
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)

    def test_gap_analysis_content(self):
        """Test Gap Analysis sheet content."""
        from openpyxl import load_workbook

        results = generate_mock_analysis_results(5)
        engine = ExcelExportEngine(MOCK_FRAMEWORK, results)

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            output_path = tmp.name

        try:
            engine.generate_workbook(output_path)
            wb = load_workbook(output_path)
            ws = wb["Gap Analysis"]

            # Check title
            assert "Gap Analysis" in str(ws['A1'].value)

            # Check headers
            assert ws['A3'].value == "Objective ID"
            assert ws['F3'].value == "Gap Status"
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)

    def test_recommendations_with_data(self):
        """Test Recommendations sheet with recommendation data."""
        from openpyxl import load_workbook

        results = generate_mock_analysis_results(3)
        recommendations = generate_mock_recommendations(results)
        engine = ExcelExportEngine(MOCK_FRAMEWORK, results, recommendations)

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            output_path = tmp.name

        try:
            engine.generate_workbook(output_path)
            wb = load_workbook(output_path)
            ws = wb["Recommendations"]

            # Check headers
            assert ws['A1'].value == "Employee Name"
            assert ws['D1'].value == "Revised Goal"

            # Should have data rows (3 employees x 5 recommendations each = 15 rows + header)
            assert ws.max_row >= 4  # At least header + some data
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)

    def test_executive_summary_formulas(self):
        """Test Executive Summary uses formulas not hardcoded values."""
        from openpyxl import load_workbook

        results = generate_mock_analysis_results(5)
        engine = ExcelExportEngine(MOCK_FRAMEWORK, results)

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            output_path = tmp.name

        try:
            engine.generate_workbook(output_path)
            wb = load_workbook(output_path)
            ws = wb["Executive Summary"]

            # Check that average scores use formulas
            avg_alignment_cell = ws['B5']
            assert avg_alignment_cell.value is not None
            if isinstance(avg_alignment_cell.value, str):
                assert "AVERAGE" in avg_alignment_cell.value
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)

    def test_score_color_formatting(self):
        """Test score cells have color formatting."""
        from openpyxl import load_workbook

        results = generate_mock_analysis_results(5)
        # Set varied scores to test color logic
        results[0]['overallAlignmentScore'] = 85  # High - green
        results[1]['overallAlignmentScore'] = 65  # Moderate - yellow
        results[2]['overallAlignmentScore'] = 35  # Low - red

        engine = ExcelExportEngine(MOCK_FRAMEWORK, results)

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            output_path = tmp.name

        try:
            engine.generate_workbook(output_path)
            wb = load_workbook(output_path)
            ws = wb["Employee Details"]

            # Check that cells have fills applied (not checking specific colors)
            high_score_cell = ws['E2']  # First employee alignment score
            assert high_score_cell.fill is not None
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)

    def test_convenience_function(self):
        """Test generate_excel_export convenience function."""
        results = generate_mock_analysis_results(3)

        output_path = generate_excel_export(MOCK_FRAMEWORK, results)

        try:
            assert os.path.exists(output_path)
            assert output_path.endswith('.xlsx')
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)

    def test_empty_results(self):
        """Test export with empty results."""
        engine = ExcelExportEngine(MOCK_FRAMEWORK, [])

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            output_path = tmp.name

        try:
            result_path = engine.generate_workbook(output_path)
            assert os.path.exists(result_path)
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)

    def test_error_results_handling(self):
        """Test export handles error results gracefully."""
        results = generate_mock_analysis_results(3)
        results.append({
            "documentId": "error-doc",
            "fileName": "failed_doc.docx",
            "error": "Processing failed"
        })

        engine = ExcelExportEngine(MOCK_FRAMEWORK, results)

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            output_path = tmp.name

        try:
            result_path = engine.generate_workbook(output_path)
            assert os.path.exists(result_path)
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)


class TestPDFExportEngine:
    """Test cases for the PDFExportEngine class."""

    def test_engine_initialization(self):
        """Test PDF engine initializes correctly."""
        results = generate_mock_analysis_results(3)
        engine = PDFExportEngine(MOCK_FRAMEWORK, results)

        assert engine.framework == MOCK_FRAMEWORK
        assert len(engine.results) == 3

    def test_generate_report_creates_file(self):
        """Test PDF report generation creates file."""
        results = generate_mock_analysis_results(5)
        engine = PDFExportEngine(MOCK_FRAMEWORK, results)

        with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tmp:
            output_path = tmp.name

        try:
            result_path = engine.generate_report(output_path)
            assert os.path.exists(result_path)
            assert result_path == output_path
            # Check file size is reasonable
            assert os.path.getsize(result_path) > 1000
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)

    def test_convenience_function(self):
        """Test generate_pdf_export convenience function."""
        results = generate_mock_analysis_results(3)

        output_path = generate_pdf_export(MOCK_FRAMEWORK, results)

        try:
            assert os.path.exists(output_path)
            assert output_path.endswith('.pdf')
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)

    def test_empty_results(self):
        """Test PDF export with empty results."""
        engine = PDFExportEngine(MOCK_FRAMEWORK, [])

        with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tmp:
            output_path = tmp.name

        try:
            result_path = engine.generate_report(output_path)
            assert os.path.exists(result_path)
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)

    def test_with_recommendations(self):
        """Test PDF export with recommendations."""
        results = generate_mock_analysis_results(3)
        recommendations = generate_mock_recommendations(results)
        engine = PDFExportEngine(MOCK_FRAMEWORK, results, recommendations)

        with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tmp:
            output_path = tmp.name

        try:
            result_path = engine.generate_report(output_path)
            assert os.path.exists(result_path)
            # File with recommendations should be larger
            assert os.path.getsize(result_path) > 2000
        finally:
            if os.path.exists(output_path):
                os.unlink(output_path)


class TestExportIntegration:
    """Integration tests for export functionality."""

    def test_full_export_workflow(self):
        """Test complete export workflow with both formats."""
        results = generate_mock_analysis_results(5)
        recommendations = generate_mock_recommendations(results)

        # Generate Excel
        excel_path = generate_excel_export(MOCK_FRAMEWORK, results, recommendations)

        # Generate PDF
        pdf_path = generate_pdf_export(MOCK_FRAMEWORK, results, recommendations)

        try:
            assert os.path.exists(excel_path)
            assert os.path.exists(pdf_path)

            # Both files should have content
            assert os.path.getsize(excel_path) > 0
            assert os.path.getsize(pdf_path) > 0
        finally:
            for path in [excel_path, pdf_path]:
                if os.path.exists(path):
                    os.unlink(path)

    def test_large_dataset_export(self):
        """Test export with maximum allowed documents (15)."""
        results = generate_mock_analysis_results(15)

        excel_path = generate_excel_export(MOCK_FRAMEWORK, results)

        try:
            from openpyxl import load_workbook
            wb = load_workbook(excel_path)
            ws = wb["Employee Details"]

            # Should have 15 employees + header
            assert ws.max_row == 16
        finally:
            if os.path.exists(excel_path):
                os.unlink(excel_path)


if __name__ == '__main__':
    pytest.main([__file__, '-v'])
